import io

from docx import Document

from app.docx_convert import docx_to_markdown, markdown_to_docx


def test_heading_round_trip():
    data = markdown_to_docx("## 見出し")
    doc = Document(io.BytesIO(data))
    assert len(doc.paragraphs) == 1
    p = doc.paragraphs[0]
    assert p.style.name.startswith("Heading")
    assert p.text == "見出し"
    assert docx_to_markdown(data) == "## 見出し"


def test_blockquote_round_trip():
    data = markdown_to_docx("> 引用文")
    doc = Document(io.BytesIO(data))
    p = doc.paragraphs[0]
    assert "Quote" in p.style.name
    assert p.text == "引用文"
    assert docx_to_markdown(data) == "> 引用文"


def test_bold_italic_runs():
    data = markdown_to_docx("普通 **太字** 普通 *斜体* 普通 ***太字斜体***")
    doc = Document(io.BytesIO(data))
    p = doc.paragraphs[0]
    bold_runs = [r for r in p.runs if r.bold and not r.italic]
    italic_runs = [r for r in p.runs if r.italic and not r.bold]
    bold_italic_runs = [r for r in p.runs if r.bold and r.italic]
    assert any(r.text == "太字" for r in bold_runs)
    assert any(r.text == "斜体" for r in italic_runs)
    assert any(r.text == "太字斜体" for r in bold_italic_runs)


def test_bold_italic_markdown_to_text_round_trip():
    original = "本文 **太字** さらに *斜体* 文章"
    data = markdown_to_docx(original)
    assert docx_to_markdown(data) == original


def test_multiple_paragraphs_separated_by_blank_lines():
    content = "段落1\n\n段落2"
    data = markdown_to_docx(content)
    assert docx_to_markdown(data) == content


def test_plain_paragraph_has_no_bold_or_italic():
    data = markdown_to_docx("ふつうの文章です")
    doc = Document(io.BytesIO(data))
    p = doc.paragraphs[0]
    assert p.text == "ふつうの文章です"
    assert not p.style.name.startswith("Heading")
    assert "Quote" not in p.style.name
